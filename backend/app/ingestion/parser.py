"""Document parsers: PyMuPDF (PDF, with table detection and Tesseract OCR), python-docx (DOCX), plain Python (TXT/MD).

Each parser yields :class:`TextBlock` objects carrying page and section information so that
chunks - and therefore citations - keep precise provenance. Tables are emitted as single
Markdown blocks (``kind="table"``) so the chunker never splits a row from its header.
"""

from __future__ import annotations

import io
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.core.errors import DocumentProcessingError, InvalidFileError
from app.utils.text import clean_text


@dataclass
class TextBlock:
    text: str
    page_number: int | None = None
    section: str | None = None
    kind: str = "text"  # "text" | "table"


@dataclass(frozen=True)
class ParseOptions:
    ocr_enabled: bool = True
    ocr_language: str = "eng"
    ocr_dpi: int = 300
    ocr_min_page_chars: int = 25


def markdown_table(rows: list[list[str]]) -> str:
    """Render rows (first row = header) as a GitHub-style Markdown table."""
    width = max(len(r) for r in rows)
    norm = [[" ".join(str(c or "").split()).replace("|", "/") for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(norm[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(lines)


def ocr_available() -> bool:
    import shutil

    return shutil.which("tesseract") is not None


@dataclass
class ParsedDocument:
    blocks: list[TextBlock]
    title: str | None
    page_count: int | None
    metadata: dict[str, object] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks)


_NUMBERED_HEADING = re.compile(r"^(\d+(\.\d+)*\.?|[IVX]+\.)\s+[A-Z][^.!?]{1,80}$")


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 90 or line.endswith((".", ",", ";", ":", "?", "!")):
        return bool(_NUMBERED_HEADING.match(line)) if line else False
    if _NUMBERED_HEADING.match(line):
        return True
    words = line.split()
    if len(words) > 10:
        return False
    capitalised = sum(1 for w in words if w[:1].isupper() or not w[:1].isalpha())
    return capitalised / len(words) >= 0.75 or line.isupper()


def _flush_paragraph(blocks: list[TextBlock], paragraph: list[str], page: int, section: str | None) -> None:
    if paragraph:
        text = clean_text(" ".join(paragraph))
        if text:
            blocks.append(TextBlock(text=text, page_number=page, section=section))
        paragraph.clear()


def _join_spans(spans: list[dict[str, Any]]) -> str:
    """Join spans, inserting a space where a visual gap separates words (OCR emits one span per word)."""
    out = ""
    prev: dict[str, Any] | None = None
    for span in spans:
        text = span["text"]
        if prev is not None and not out.endswith(" ") and not text.startswith(" "):
            gap = span["bbox"][0] - prev["bbox"][2]
            if gap > 0.15 * float(span.get("size") or 10):
                out += " "
        out += text
        prev = span
    return out.strip()


def _page_tables(page: Any) -> list[tuple[float, float, float, float, str]]:
    """Ruled tables on a page as (x0, y0, x1, y1, markdown); failures degrade to plain text."""
    try:
        found = page.find_tables()
    except Exception:
        return []
    out = []
    for table in found.tables:
        rows = [row for row in table.extract() if any(c for c in row)]
        if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
            out.append((*table.bbox, markdown_table(rows)))
    return out


def _inside(bbox: tuple[float, ...], table: tuple[float, ...]) -> bool:
    x0, y0, x1, y1 = bbox[:4]
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    return table[0] - 2 <= cx <= table[2] + 2 and table[1] - 2 <= cy <= table[3] + 2


def parse_pdf(data: bytes, options: ParseOptions | None = None) -> ParsedDocument:
    import pymupdf

    options = options or ParseOptions()
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise InvalidFileError("Unable to open PDF") from exc
    if doc.needs_pass:
        raise InvalidFileError("Encrypted PDFs are not supported")
    blocks: list[TextBlock] = []
    meta_title = (doc.metadata or {}).get("title") or None
    first_heading: str | None = None
    section: str | None = None
    ocr_pages: list[int] = []
    table_count = 0
    can_ocr = options.ocr_enabled and ocr_available()
    try:
        # Pass 1: text (OCR for image-only pages), tables, and font sizes for heading detection.
        sizes: list[float] = []
        pages: list[tuple[dict[str, Any], list[tuple[float, float, float, float, str]]]] = []
        for index, page in enumerate(doc, start=1):
            textpage = None
            if can_ocr and len(page.get_text().strip()) < options.ocr_min_page_chars:
                try:
                    textpage = page.get_textpage_ocr(language=options.ocr_language, dpi=options.ocr_dpi, full=True)
                    ocr_pages.append(index)
                except Exception:
                    textpage = None
            info = page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT, textpage=textpage)
            tables = [] if textpage is not None else _page_tables(page)
            pages.append((info, tables))
            for blk in info.get("blocks", []):
                if any(_inside(blk["bbox"], t) for t in tables):
                    continue
                for line in blk.get("lines", []):
                    for span in line.get("spans", []):
                        if span.get("text", "").strip():
                            sizes.append(round(float(span.get("size", 0)), 1))
        body_size = statistics.median(sizes) if sizes else 0.0
        # Pass 2: blocks in reading order, tables inserted where they appear.
        for page_index, (info, tables) in enumerate(pages, start=1):
            pending = sorted(tables, key=lambda t: t[1])
            paragraph: list[str] = []
            for blk in info.get("blocks", []):
                while pending and pending[0][1] <= blk["bbox"][1]:
                    _flush_paragraph(blocks, paragraph, page_index, section)
                    blocks.append(TextBlock(pending.pop(0)[4], page_index, section, kind="table"))
                    table_count += 1
                if any(_inside(blk["bbox"], t) for t in tables):
                    continue
                for line in blk.get("lines", []):
                    spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                    if not spans:
                        continue
                    line_text = _join_spans(spans)
                    max_size = max(float(s.get("size", 0)) for s in spans)
                    bold = any(int(s.get("flags", 0)) & 16 for s in spans)
                    is_heading = (
                        len(line_text) <= 90
                        and (max_size >= body_size * 1.15 or (bold and _looks_like_heading(line_text)))
                        and not line_text.endswith((".", ","))
                    )
                    if is_heading:
                        _flush_paragraph(blocks, paragraph, page_index, section)
                        section = clean_text(line_text)
                        first_heading = first_heading or section
                    else:
                        paragraph.append(line_text)
                _flush_paragraph(blocks, paragraph, page_index, section)
            for table in pending:
                blocks.append(TextBlock(table[4], page_index, section, kind="table"))
                table_count += 1
    finally:
        page_count = doc.page_count
        doc.close()
    if not blocks:
        hint = "OCR found no text" if can_ocr else "scanned PDFs need OCR - install Tesseract or enable OCR_ENABLED"
        raise DocumentProcessingError(f"No extractable text found in PDF ({hint})")
    meta: dict[str, object] = {"parser": "pymupdf", "tables": table_count, "ocr_pages": len(ocr_pages)}
    if ocr_pages:
        meta["ocr_language"] = options.ocr_language
    return ParsedDocument(blocks, meta_title or first_heading, page_count, meta)


def parse_docx(data: bytes) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise InvalidFileError("Unable to open DOCX document") from exc
    blocks: list[TextBlock] = []
    section: str | None = None
    title = document.core_properties.title or None
    first_heading: str | None = None
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    tables = 0
    # Walk the body in order so tables stay next to the text that introduces them.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            table = Table(child, document)
            rows = [[clean_text(c.text) for c in row.cells] for row in table.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                blocks.append(TextBlock(text=markdown_table(rows), section=section, kind="table"))
                tables += 1
            continue
        if tag != "p":
            continue
        para = Paragraph(child, document)
        text = clean_text(para.text)
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.startswith("Heading") or style == "Title":
            section = text
            first_heading = first_heading or text
            if style == "Title" and not title:
                title = text
            continue
        blocks.append(TextBlock(text=text, section=section))
    if not blocks:
        raise DocumentProcessingError("No text found in DOCX document")
    return ParsedDocument(blocks, title or first_heading, None, {"parser": "python-docx", "tables": tables})


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_INLINE = [
    (re.compile(r"!\[([^\]]*)\]\([^)]*\)"), r"\1"),
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),
    (re.compile(r"`{1,3}([^`]*)`{1,3}"), r"\1"),
    (re.compile(r"(\*\*|__)(.*?)\1"), r"\2"),
    (re.compile(r"(?<!\w)[*_](.*?)[*_](?!\w)"), r"\1"),
]


def _strip_md(text: str) -> str:
    for pattern, repl in _MD_INLINE:
        text = pattern.sub(repl, text)
    return text


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def parse_markdown(data: bytes) -> ParsedDocument:
    text = _decode(data)
    blocks: list[TextBlock] = []
    section: str | None = None
    title: str | None = None
    paragraph: list[str] = []
    in_code = False

    def flush() -> None:
        if not paragraph:
            return
        lines = [ln.strip() for ln in paragraph]
        if len(lines) >= 2 and all(ln.startswith("|") for ln in lines) and re.match(r"^\|[\s:\-|]+\|$", lines[1]):
            rows = [[_strip_md(c.strip()) for c in ln.strip("|").split("|")] for i, ln in enumerate(lines) if i != 1]
            blocks.append(TextBlock(text=markdown_table(rows), section=section, kind="table"))
        else:
            cleaned = clean_text(_strip_md("\n".join(paragraph)))
            if cleaned:
                blocks.append(TextBlock(text=cleaned, section=section))
        paragraph.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        match = _MD_HEADING.match(line) if not in_code else None
        if match:
            flush()
            section = _strip_md(match.group(2)).strip()
            if len(match.group(1)) == 1 and title is None:
                title = section
        elif not line.strip():
            flush()
        else:
            paragraph.append(line.rstrip())
    flush()
    if not blocks:
        raise DocumentProcessingError("Markdown document contains no text")
    return ParsedDocument(blocks, title, None, {"parser": "markdown"})


def parse_text(data: bytes) -> ParsedDocument:
    text = clean_text(_decode(data))
    blocks: list[TextBlock] = []
    section: str | None = None
    title: str | None = None
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        lines = para.split("\n")
        if len(lines) == 1 and _looks_like_heading(lines[0]) and len(lines[0].split()) <= 8:
            section = lines[0].strip()
            title = title or section
            continue
        if _looks_like_heading(lines[0]) and len(lines) > 1 and len(lines[0].split()) <= 8:
            section = lines[0].strip()
            title = title or section
            para = "\n".join(lines[1:])
        blocks.append(TextBlock(text=" ".join(para.split()), section=section))
    if not blocks:
        raise DocumentProcessingError("Text document is empty")
    return ParsedDocument(blocks, title, None, {"parser": "text"})


PARSERS = {"pdf": parse_pdf, "docx": parse_docx, "md": parse_markdown, "txt": parse_text}


def parse_document(data: bytes, file_type: str, filename: str, options: ParseOptions | None = None) -> ParsedDocument:
    parser = PARSERS.get(file_type)
    if parser is None:
        raise InvalidFileError(f"No parser for file type {file_type}")
    parsed = parse_pdf(data, options) if file_type == "pdf" else parser(data)
    if not parsed.title:
        parsed.title = Path(filename).stem.replace("_", " ").replace("-", " ").strip() or filename
    return parsed
