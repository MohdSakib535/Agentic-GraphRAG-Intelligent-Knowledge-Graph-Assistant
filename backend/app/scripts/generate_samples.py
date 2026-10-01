"""Generate the sample enterprise corpus (PDF, DOCX, Markdown, TXT).

    python -m app.scripts.generate_samples [output_dir]
"""

from __future__ import annotations

import sys
from pathlib import Path

ARCHITECTURE_PAGES = [
    (
        "TechCorp Platform Architecture",
        [
            "This document describes the reference architecture used by TechCorp engineering teams. "
            "TechCorp is based in Bangalore and builds data-intensive products for enterprise customers.",
            "The platform is organised around event streaming, caching, relational storage and graph analytics. "
            "Each project selects components from this reference architecture.",
        ],
    ),
    (
        "Event Streaming with Kafka",
        [
            "Kafka is a distributed event streaming platform used for high-throughput, fault-tolerant data pipelines. "
            "Kafka stores events in partitioned, replicated topics so that producers and consumers are decoupled.",
            "At TechCorp, Kafka is the backbone for asynchronous communication between services. "
            "Consumer groups allow several instances of a service to share the load of a topic.",
        ],
    ),
    (
        "Project Alpha Architecture",
        [
            "Rahul manages Project Alpha. Project Alpha is the real-time order analytics platform of TechCorp.",
            "Project Alpha uses Kafka, Redis and PostgreSQL. The ingestion service of Project Alpha is built with FastAPI.",
            "Redis is used in Project Alpha as a low-latency cache for aggregated metrics, while PostgreSQL stores "
            "the curated order history.",
        ],
    ),
    (
        "Project Gamma Architecture",
        [
            "Priya manages Project Gamma. Project Gamma is the supplier knowledge graph initiative.",
            "Project Gamma uses Kafka, Neo4j and FastAPI. Project Gamma depends on Project Alpha for order events.",
            "Neo4j is a native graph database that stores data as nodes and relationships and is queried with Cypher.",
        ],
    ),
]

PROJECT_OVERVIEW = {
    "title": "Project Portfolio Overview",
    "sections": [
        (
            "Project Alpha",
            [
                "Project Alpha delivers real-time order analytics for TechCorp customers.",
                "Rahul manages Project Alpha. Amit works on Project Alpha. Neha works on Project Alpha.",
                "Project Alpha uses Kafka for event streaming and Redis for caching.",
            ],
        ),
        (
            "Project Beta",
            [
                "Project Beta is the internal customer support portal.",
                "Rahul manages Project Beta. Priya works on Project Beta.",
                "Project Beta uses Django, PostgreSQL and Redis.",
            ],
        ),
        (
            "Project Gamma",
            [
                "Project Gamma builds a supplier knowledge graph.",
                "Priya manages Project Gamma. Neha works on Project Gamma.",
                "Project Gamma uses Kafka and Neo4j.",
            ],
        ),
    ],
}

TEAM_DIRECTORY = """# TechCorp Team Directory

## Leadership

Rahul is the engineering manager at TechCorp. Rahul works for TechCorp and leads the Engineering Department.
Rahul has managed delivery programmes at TechCorp since 2019.

## Engineers

Amit, a senior backend developer, works for TechCorp. Amit reports to Rahul.
Amit uses Kafka and Redis every day to build streaming pipelines.

Priya, a lead data engineer, works for TechCorp. Priya reports to Rahul.
Priya uses Neo4j and Django.

Neha, a backend developer, works for TechCorp. Neha reports to Rahul.
Neha uses FastAPI and Kafka.

## Teams

The Data Platform Team belongs to the Engineering Department.
The Engineering Department is based in Bangalore.
"""

GLOSSARY = """Technology Glossary

Kafka
Kafka is a distributed event streaming platform originally developed at LinkedIn and now maintained by the Apache Software Foundation. It is used to build real-time data pipelines and streaming applications.

Redis
Redis is an in-memory key-value data store commonly used as a cache, message broker and session store. Redis offers sub-millisecond latency.

PostgreSQL
PostgreSQL is an open-source object-relational database system known for reliability, SQL compliance and extensibility.

FastAPI
FastAPI is a modern Python web framework for building APIs with automatic OpenAPI documentation based on type hints.

Django
Django is a high-level Python web framework that encourages rapid development with an included ORM and admin interface.

Neo4j
Neo4j is a graph database management system. Neo4j stores nodes and relationships natively and uses the Cypher query language.
"""


def write_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    for heading, paragraphs in ARCHITECTURE_PAGES:
        page = doc.new_page()
        y = 72
        page.insert_text((72, y), heading, fontsize=18, fontname="helv")
        y += 36
        for para in paragraphs:
            rect = pymupdf.Rect(72, y, 540, y + 160)
            page.insert_textbox(rect, para, fontsize=11, fontname="helv")
            y += 20 * (len(para) // 85 + 2)
    doc.set_metadata({"title": "TechCorp Platform Architecture", "author": "TechCorp Engineering"})
    doc.save(path)
    doc.close()


def write_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.core_properties.title = PROJECT_OVERVIEW["title"]
    document.add_heading(PROJECT_OVERVIEW["title"], level=0)
    for heading, paragraphs in PROJECT_OVERVIEW["sections"]:
        document.add_heading(heading, level=1)
        for para in paragraphs:
            document.add_paragraph(para)
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Project", "Manager", "Status"
    for project, manager, status in (("Project Alpha", "Rahul", "Active"), ("Project Beta", "Rahul", "Active"),
                                     ("Project Gamma", "Priya", "Pilot")):
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = project, manager, status
    document.save(path)


def main(out_dir: str | None = None) -> None:
    out = Path(out_dir or Path(__file__).resolve().parents[2] / "data" / "samples")
    out.mkdir(parents=True, exist_ok=True)
    write_pdf(out / "architecture.pdf")
    write_docx(out / "project-overview.docx")
    (out / "team-directory.md").write_text(TEAM_DIRECTORY)
    (out / "technology-glossary.txt").write_text(GLOSSARY)
    print(f"Sample documents written to {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
